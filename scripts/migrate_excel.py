#!/usr/bin/env python3
"""money-lens: 家計簿Excel(収支表.xlsx) -> money-lens-data JSON 移行スクリプト

docs/data-migration.md の仕様に厳密に従う。

使い方:
    # 全件移行(初回用)。transactions/*.json と months.json も上書きするので注意
    python3 scripts/migrate_excel.py --input .claude/tmp/budget.xlsx --output ../money-lens-data

    # 給与のみ追記(移行後に給与行を追加する運用)。既存レコードは変更せず、未登録idのみ追記。
    # transactions / months は一切書き出さない
    python3 scripts/migrate_excel.py --input .claude/tmp/budget.xlsx --output ../money-lens-data \\
        --salaries-only [--until YYYY-MM] [--dry-run]
    #   --until YYYY-MM : その年月より後の行を除外 (--salaries-only 専用)
    #   --dry-run       : 追記内容を表示するだけで書き出さない (--salaries-only 専用)

出力:
    <output>/transactions/2025.json
    <output>/transactions/2026.json
    <output>/salaries.json
    <output>/months.json

categories.json / budgets.json / settings.json はこのスクリプトでは生成しない
(data-model.md の初期値を手動配置する運用)。ただし検証時に <output>/categories.json
が存在すればそれを読み込み、未知カテゴリの警告に使う。
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# 収支一覧シートの列番号 (1-indexed)
# ---------------------------------------------------------------------------
TX_SHEET = "収支一覧"
TX_HEADER_ROW = 2
TX_DATA_START_ROW = 3
COL_TX_SEQ = 2  # '#'
COL_TX_DATE = 3
COL_TX_AMOUNT = 4
COL_TX_CURRENCY = 5
COL_TX_DESCRIPTION = 6
COL_TX_MEMO = 7
COL_TX_RECEIPT = 8
COL_TX_CATEGORY = 9
COL_TX_ACCOUNT_NAME = 10
COL_TX_ACCOUNT_NUMBER = 11  # 保存しない
COL_TX_WORK = 12

# ---------------------------------------------------------------------------
# 給与一覧シートの列定義
# 列番号はハードコードせず、3行目のヘッダ名から解決する (Excel側で列が挿入されても追従するため)。
# '#'(列2)と等級(列3, ヘッダ無し)のみ位置固定。
# ---------------------------------------------------------------------------
SALARY_SHEET = "給与一覧"
SALARY_HEADER_ROW = 3
SALARY_DATA_START_ROW = 4
COL_S_SEQ = 2
COL_S_GRADE = 3

# 単一値フィールド: 内部名 -> ヘッダ名
SALARY_SCALAR_HEADERS = {
    "year": "年",
    "month": "月",
    "type": "給与/賞与",
    "grossTotal": "支給合計",
    "deductionTotal": "控除合計",
    "netPay": "差引支給額",
    "bankTransfer": "銀行振込額",
    "taxableAmount": "課税対象額",
}

PAYMENT_FIELDS = [
    "基本給",
    "住宅手当",
    "前払退職金",
    "時間外勤務手当",
    "深夜手当",
    "賞与基本給",
    "特別支給",
    "その他支給",
    "基本給調整",
]

DEDUCTION_FIELDS = [
    "健康保険",
    "厚生年金",
    "雇用保険",
    "社会保険料合計",
    "所得税",
    "住民税",
    "年調過不足額",
    "その他控除額",
    "株式購入",
    "レインバース",
]

ATTENDANCE_FIELDS = [
    "所定勤務日数",
    "勤務日数",
    "欠勤日数",
    "時間外勤務時間",
    "深夜時間",
    "休日時間",
    "精算残業時間",
    "60時間超",
]

# 以下のヘッダ列は存在しても読み取らず、保存もしない (JSONスキーマに対応フィールドが無い):
#   「出勤日調整」(支給グループ) / 「子育て支援金」(控除グループ)
#   「第1口座」: 実データでは銀行振込額と常に同値の重複列

# ---------------------------------------------------------------------------
# 統計シートの列番号 (1-indexed)
# ---------------------------------------------------------------------------
MONTHS_SHEET = "統計"
MONTHS_DATA_START_ROW = 4
COL_M_BALANCE_START = 4  # 月初金額
COL_M_BALANCE_END = 5  # 月末金額
COL_M_YEAR = 6  # 年
COL_M_MONTH = 11  # 月 ('01'..'12')
COL_M_STAR = 16  # ★
COL_M_NOTE = 17  # メモ

STAR_MARK = "★"  # ★


def num(v):
    """Excelの数値セルを int/float/None に正規化する。整数値のfloatはintへ丸める。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        if v.is_integer():
            return int(v)
        return v
    return v


def money(v):
    """金額(円)は必ず整数として扱う。整数化できない場合はエラーにする。"""
    if v is None:
        return None
    if isinstance(v, float):
        if not v.is_integer():
            raise ValueError(f"金額フィールドが整数でない: {v!r}")
        return int(v)
    if isinstance(v, int):
        return v
    raise ValueError(f"金額フィールドが数値でない: {v!r}")


def strp(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def load_categories(output_dir: Path):
    cat_path = output_dir / "categories.json"
    if not cat_path.exists():
        return None
    with cat_path.open("r", encoding="utf-8") as f:
        data = json.load(f)
    return set(data.get("categories", {}).keys())


# ---------------------------------------------------------------------------
# transactions
# ---------------------------------------------------------------------------
def parse_transactions(ws):
    raw_rows = []
    r = TX_DATA_START_ROW
    while True:
        date_val = ws.cell(row=r, column=COL_TX_DATE).value
        if date_val is None:
            break
        raw_rows.append(
            {
                "date": strp(date_val),
                "amount": money(ws.cell(row=r, column=COL_TX_AMOUNT).value),
                "currency": strp(ws.cell(row=r, column=COL_TX_CURRENCY).value) or "JPY",
                "description": strp(ws.cell(row=r, column=COL_TX_DESCRIPTION).value) or "",
                "memo": strp(ws.cell(row=r, column=COL_TX_MEMO).value),
                "category": strp(ws.cell(row=r, column=COL_TX_CATEGORY).value) or "Uncategorized",
                "account": strp(ws.cell(row=r, column=COL_TX_ACCOUNT_NAME).value) or "",
                "work": strp(ws.cell(row=r, column=COL_TX_WORK).value) or "",
                "_row": r,
            }
        )
        r += 1

    by_year = defaultdict(list)
    for row in raw_rows:
        year = int(row["date"][:4])
        by_year[year].append(row)

    transactions_by_year = {}
    for year, rows in by_year.items():
        # 日付順(同日内は元の行順を維持する安定ソート)に並べてから連番採番
        rows_sorted = sorted(rows, key=lambda x: (x["date"], x["_row"]))
        out = []
        for i, row in enumerate(rows_sorted, start=1):
            out.append(
                {
                    "id": f"t-{year}-{i:06d}",
                    "date": row["date"],
                    "amount": row["amount"],
                    "currency": row["currency"],
                    "description": row["description"],
                    "memo": row["memo"],
                    "category": row["category"],
                    "account": row["account"],
                    "work": row["work"],
                    "isSpecial": False,
                    "specialNote": None,
                }
            )
        transactions_by_year[year] = out

    return transactions_by_year, raw_rows


# ---------------------------------------------------------------------------
# salaries
# ---------------------------------------------------------------------------
def resolve_salary_columns(ws):
    """3行目のヘッダ名 -> 列番号(1-indexed)。必要なヘッダが欠けていれば終了コード1で終了。"""
    header_to_col = {}
    for c in range(1, ws.max_column + 1):
        h = strp(ws.cell(row=SALARY_HEADER_ROW, column=c).value)
        if h is not None and h not in header_to_col:
            header_to_col[h] = c
    required = (
        list(SALARY_SCALAR_HEADERS.values())
        + PAYMENT_FIELDS
        + DEDUCTION_FIELDS
        + ATTENDANCE_FIELDS
    )
    missing = [h for h in required if h not in header_to_col]
    if missing:
        print(
            f"給与一覧の{SALARY_HEADER_ROW}行目に必要なヘッダが見つかりません: {missing}",
            file=sys.stderr,
        )
        sys.exit(1)
    return header_to_col


def parse_salaries(ws):
    cols = resolve_salary_columns(ws)
    col = {key: cols[h] for key, h in SALARY_SCALAR_HEADERS.items()}
    raw_rows = []
    r = SALARY_DATA_START_ROW
    while r <= ws.max_row:
        gross = ws.cell(row=r, column=col["grossTotal"]).value
        year_val = ws.cell(row=r, column=col["year"]).value
        month_val = ws.cell(row=r, column=col["month"]).value
        grade_val = ws.cell(row=r, column=COL_S_GRADE).value
        # 年/月/給与・賞与区分は未到来の将来月にも下書きとして入力済みのため、
        # それだけでは実データ行かどうか判定できない。給与体系(grade)が入力されている
        # 行のみを実データとして扱う(金額が未記入の欠損レコードでもgradeがあれば含める)。
        if year_val is None or month_val is None or grade_val is None:
            r += 1
            continue

        year = int(str(year_val).strip().rstrip("年"))
        month = int(str(month_val).strip().rstrip("月"))
        s_type = strp(ws.cell(row=r, column=col["type"]).value)

        payments = {
            name: money(ws.cell(row=r, column=cols[name]).value) for name in PAYMENT_FIELDS
        }
        deductions = {
            name: money(ws.cell(row=r, column=cols[name]).value) for name in DEDUCTION_FIELDS
        }
        attendance = {
            name: num(ws.cell(row=r, column=cols[name]).value) for name in ATTENDANCE_FIELDS
        }

        raw_rows.append(
            {
                "year": year,
                "month": month,
                "type": s_type,
                "grade": strp(ws.cell(row=r, column=COL_S_GRADE).value),
                "grossTotal": money(gross),
                "deductionTotal": money(ws.cell(row=r, column=col["deductionTotal"]).value),
                "netPay": money(ws.cell(row=r, column=col["netPay"]).value),
                "payments": payments,
                "deductions": deductions,
                "attendance": attendance,
                "bankTransfer": money(ws.cell(row=r, column=col["bankTransfer"]).value),
                "taxableAmount": money(ws.cell(row=r, column=col["taxableAmount"]).value),
                "note": None,
                "_row": r,
            }
        )
        r += 1

    # 同年月に給与と賞与がある場合、賞与の id に -bonus サフィックスを付与
    groups = defaultdict(list)
    for row in raw_rows:
        groups[(row["year"], row["month"])].append(row)

    salaries = []
    for row in raw_rows:
        key = (row["year"], row["month"])
        has_salary_sibling = any(
            sib["type"] == "給与" for sib in groups[key] if sib is not row
        )
        base_id = f"s-{row['year']}-{row['month']:02d}"
        if row["type"] == "賞与" and has_salary_sibling:
            sid = f"{base_id}-bonus"
        else:
            sid = base_id
        salaries.append(
            {
                "id": sid,
                "year": row["year"],
                "month": row["month"],
                "type": row["type"],
                "grade": row["grade"],
                "grossTotal": row["grossTotal"],
                "deductionTotal": row["deductionTotal"],
                "netPay": row["netPay"],
                "payments": row["payments"],
                "deductions": row["deductions"],
                "attendance": row["attendance"],
                "bankTransfer": row["bankTransfer"],
                "taxableAmount": row["taxableAmount"],
                "note": row["note"],
            }
        )

    return salaries, raw_rows


# ---------------------------------------------------------------------------
# months
# ---------------------------------------------------------------------------
def parse_months(ws):
    months = []
    r = MONTHS_DATA_START_ROW
    while r <= ws.max_row:
        balance_start = ws.cell(row=r, column=COL_M_BALANCE_START).value
        balance_end = ws.cell(row=r, column=COL_M_BALANCE_END).value
        year_val = ws.cell(row=r, column=COL_M_YEAR).value
        month_val = ws.cell(row=r, column=COL_M_MONTH).value

        if balance_start is None and balance_end is None:
            # データ未記入の将来月はスキップ
            r += 1
            continue
        if year_val is None or month_val is None:
            r += 1
            continue

        star_val = ws.cell(row=r, column=COL_M_STAR).value
        note_val = strp(ws.cell(row=r, column=COL_M_NOTE).value)
        is_special = bool(star_val and str(star_val).strip())

        month_str = str(month_val).strip()
        if month_str.isdigit():
            month_str = f"{int(month_str):02d}"

        months.append(
            {
                "month": f"{int(year_val)}-{month_str}",
                "isSpecial": is_special,
                "note": note_val,
                "balanceStart": money(balance_start),
                "balanceEnd": money(balance_end),
            }
        )
        r += 1

    return months


# ---------------------------------------------------------------------------
# 検証
# ---------------------------------------------------------------------------
def check_netpay(salaries):
    """netPay == grossTotal - deductionTotal の不一致 [(id, expected, actual)] を返す。"""
    mismatches = []
    for s in salaries:
        if s["grossTotal"] is None or s["deductionTotal"] is None:
            # 金額未記入の欠損レコード(例: s-2023-12-bonus)。null同士なら整合とみなす。
            if s["netPay"] is not None:
                mismatches.append((s["id"], None, s["netPay"]))
            continue
        expected = s["grossTotal"] - s["deductionTotal"]
        if expected != s["netPay"]:
            mismatches.append((s["id"], expected, s["netPay"]))
    return mismatches


def run_validation(ws_tx, transactions_by_year, raw_tx_rows, salaries, categories_known):
    checks = []

    # 1. 取引件数一致
    excel_count = len(raw_tx_rows)
    json_count = sum(len(v) for v in transactions_by_year.values())
    checks.append(
        (
            f"取引件数一致 (Excel={excel_count} / JSON={json_count})",
            excel_count == json_count,
        )
    )

    # 2. 年別件数一致
    excel_year_counts = Counter(int(row["date"][:4]) for row in raw_tx_rows)
    json_year_counts = {y: len(v) for y, v in transactions_by_year.items()}
    year_ok = dict(excel_year_counts) == json_year_counts
    checks.append(
        (
            f"年別件数一致 (Excel={dict(excel_year_counts)} / JSON={json_year_counts})",
            year_ok,
        )
    )

    # 2b. カテゴリ別件数一致(年をまたいだ全体で比較)
    excel_cat_counts = Counter(row["category"] for row in raw_tx_rows)
    json_cat_counts = Counter()
    for txs in transactions_by_year.values():
        for t in txs:
            json_cat_counts[t["category"]] += 1
    cat_ok = excel_cat_counts == json_cat_counts
    checks.append((f"カテゴリ別件数一致 (種類数 Excel={len(excel_cat_counts)} / JSON={len(json_cat_counts)})", cat_ok))

    # 3. 年ごとの金額合計一致
    excel_year_sum = defaultdict(int)
    for row in raw_tx_rows:
        excel_year_sum[int(row["date"][:4])] += row["amount"]
    sum_ok = True
    sum_detail = []
    for year, txs in transactions_by_year.items():
        json_sum = sum(t["amount"] for t in txs)
        ok = json_sum == excel_year_sum[year]
        sum_detail.append(f"{year}: Excel={excel_year_sum[year]} JSON={json_sum}")
        sum_ok = sum_ok and ok
    checks.append((f"年ごとの金額合計一致 ({'; '.join(sum_detail)})", sum_ok))

    # 4. 給与件数 + netPay整合
    salary_count_ok = len(salaries) == 56
    checks.append((f"給与件数56件 (実際={len(salaries)})", salary_count_ok))

    netpay_mismatches = check_netpay(salaries)
    netpay_ok = len(netpay_mismatches) == 0
    detail = "" if netpay_ok else f" 不一致: {netpay_mismatches}"
    checks.append((f"netPay = grossTotal - deductionTotal 整合{detail}", netpay_ok))

    # 5. 未知カテゴリ警告
    unknown = []
    if categories_known is not None:
        used_categories = set(excel_cat_counts.keys())
        unknown = sorted(used_categories - categories_known)
    unknown_ok = len(unknown) == 0
    if categories_known is None:
        checks.append(("未知カテゴリチェック (categories.json が見つからずスキップ)", True))
    else:
        checks.append((f"未知カテゴリなし (未知={unknown})", unknown_ok))

    return checks


def parse_year_month(text):
    try:
        y, m = text.split("-")
        if len(text) != 7 or len(m) != 2:
            raise ValueError
        y, m = int(y), int(m)
        if not 1 <= m <= 12:
            raise ValueError
    except ValueError:
        print(f"--until は YYYY-MM 形式で指定してください: {text}", file=sys.stderr)
        sys.exit(1)
    return (y, m)


def salary_sort_key(s):
    return (s["year"], s["month"], 0 if s["type"] == "給与" else 1)


def run_salaries_only(wb, output_dir, until, dry_run):
    """給与一覧のみ解析し、salaries.json に未登録idのレコードだけを追記する。"""
    salaries, _ = parse_salaries(wb[SALARY_SHEET])
    if until is not None:
        # id採番(-bonus判定)は全行で済んでいるので、その後で絞り込む
        salaries = [s for s in salaries if (s["year"], s["month"]) <= until]

    salaries_path = output_dir / "salaries.json"
    if not salaries_path.exists():
        print(f"既存の salaries.json が見つかりません: {salaries_path}", file=sys.stderr)
        sys.exit(1)
    with salaries_path.open("r", encoding="utf-8") as f:
        existing = json.load(f)["salaries"]
    existing_by_id = {s["id"]: s for s in existing}

    new_records = []
    for s in salaries:
        old = existing_by_id.get(s["id"])
        if old is None:
            new_records.append(s)
        elif old != s:
            keys = sorted(k for k in set(old) | set(s) if old.get(k) != s.get(k))
            print(f"WARNING: 既存レコード {s['id']} がExcelと異なります (上書きしません): {keys}")

    mismatches = check_netpay(new_records)
    for sid, expected, actual in mismatches:
        print(f"[FAIL] {sid}: netPay={actual} だが grossTotal-deductionTotal={expected}")
    if mismatches:
        print("検証にFAILがあるため、salaries.json は書き出しません。", file=sys.stderr)
        sys.exit(2)

    print(f"追記対象: {len(new_records)}件 " + ", ".join(s["id"] for s in new_records))
    if not new_records:
        print("追記するレコードはありません。")
        return

    merged = existing + new_records
    # 既存レコードの順序は変えない。既存が (年, 月, 給与→賞与) 順なら全体をソート、そうでなければ末尾に追記
    existing_keys = [salary_sort_key(s) for s in existing]
    if existing_keys == sorted(existing_keys):
        merged = sorted(merged, key=salary_sort_key)

    if dry_run:
        print(f"[dry-run] 書き出しません (書き出した場合は計 {len(merged)}件)")
        return

    with salaries_path.open("w", encoding="utf-8") as f:
        json.dump({"salaries": merged}, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"書き出し: {salaries_path} ({len(merged)}件、追記{len(new_records)}件)")


def main():
    parser = argparse.ArgumentParser(description="money-lens Excel移行スクリプト")
    parser.add_argument("--input", required=True, help="入力xlsxパス")
    parser.add_argument("--output", required=True, help="出力ディレクトリ (money-lens-dataのローカルクローン)")
    parser.add_argument("--salaries-only", action="store_true",
                        help="給与一覧のみ解析し salaries.json へ未登録idを追記する (transactions/months は触らない)")
    parser.add_argument("--until", metavar="YYYY-MM",
                        help="この年月より後の給与行を除外する (--salaries-only 専用)")
    parser.add_argument("--dry-run", action="store_true",
                        help="追記内容を表示するだけで書き出さない (--salaries-only 専用)")
    args = parser.parse_args()

    if (args.until or args.dry_run) and not args.salaries_only:
        print("--until / --dry-run は --salaries-only と併用してください", file=sys.stderr)
        sys.exit(1)
    until = parse_year_month(args.until) if args.until else None

    input_path = Path(args.input)
    output_dir = Path(args.output)
    if not input_path.exists():
        print(f"入力ファイルが見つかりません: {input_path}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(str(input_path), data_only=True)

    if args.salaries_only:
        run_salaries_only(wb, output_dir, until, args.dry_run)
        return

    ws_tx = wb[TX_SHEET]
    ws_salary = wb[SALARY_SHEET]
    ws_months = wb[MONTHS_SHEET]

    transactions_by_year, raw_tx_rows = parse_transactions(ws_tx)
    salaries, raw_salary_rows = parse_salaries(ws_salary)
    months = parse_months(ws_months)

    categories_known = load_categories(output_dir)

    checks = run_validation(ws_tx, transactions_by_year, raw_tx_rows, salaries, categories_known)

    print("=" * 70)
    print("検証結果")
    print("=" * 70)
    all_pass = True
    for label, ok in checks:
        status = "PASS" if ok else "FAIL"
        if not ok:
            all_pass = False
        print(f"[{status}] {label}")
    print("=" * 70)

    if not all_pass:
        print("検証にFAILがあるため、JSONファイルは書き出しません。原因を調査してください。", file=sys.stderr)
        sys.exit(2)

    # 書き出し
    tx_dir = output_dir / "transactions"
    tx_dir.mkdir(parents=True, exist_ok=True)
    for year, txs in sorted(transactions_by_year.items()):
        out_path = tx_dir / f"{year}.json"
        with out_path.open("w", encoding="utf-8") as f:
            json.dump({"year": year, "transactions": txs}, f, ensure_ascii=False, indent=2)
            f.write("\n")
        print(f"書き出し: {out_path} ({len(txs)}件)")

    salaries_path = output_dir / "salaries.json"
    with salaries_path.open("w", encoding="utf-8") as f:
        json.dump({"salaries": salaries}, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"書き出し: {salaries_path} ({len(salaries)}件)")

    months_path = output_dir / "months.json"
    with months_path.open("w", encoding="utf-8") as f:
        json.dump({"months": months}, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"書き出し: {months_path} ({len(months)}件)")

    print("すべての検証がPASSし、JSONを書き出しました。")


if __name__ == "__main__":
    main()
